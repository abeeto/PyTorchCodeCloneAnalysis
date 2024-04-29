from docx import Document
from win32com.client import Dispatch
path = "test.docx"
# xl=Dispatch('Excel.Application')
# xl.Visible=False
# wb=xl.Workbooks.Open(path)
# ws=wb.Sheets(1)
# rng=ws.Range("A1:B13")
# rng.CopyPicture()
# c=ws.ChartObjects().Add(0,0,rng.Width,rng.Height).Chart
# c.Activate
# c.paste()
# c.Export(r'C:\Users\FanXiaoLei\Desktop\1.png','png')
# c.Parent.Delete()
# wb.Saved=True#不保存文件
# xl.Quit()
document = Document(path)
print(document)
tables = document.tables
for table in tables:
    print("______________________________________________________")
    for i in range(1, len(table.rows)):
        result = table.cell(i, 0).text + "" + table.cell(i, 1).text + table.cell(i, 2).text + table.cell(i, 3).text
        # cell(i,0)表示第(i+1)行第1列数据，以此类推
        print(result)
